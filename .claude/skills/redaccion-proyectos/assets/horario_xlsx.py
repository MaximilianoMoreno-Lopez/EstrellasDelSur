# -*- coding: utf-8 -*-
"""Rellena la plantilla oficial KA152 Annex Timetable con el horario de un markdown.

Uso:
    py horario_xlsx.py --md horario.md --config horario.json --out Annex_Timetable.xlsx
                       [--plantilla ruta/KA152_Annex_Timetable.xlsx] [--conservar-visitas]

Markdown de entrada (en cualquier idioma):
    ## Programa del intercambio            primera sección "## ": el horario
    (texto de introducción, no se copia a la hoja)
    ### Día 1, lunes 19/04/2027. Título    un día por "### "
    Primera línea de texto: objetivo del día (va en cursiva bajo la cabecera)
    | Hora | Sesión | Objetivo | Método | Conduce |      cinco columnas, cabecera libre
    |---|---|---|---|---|
    | 9:30-11:00 | ... | ... | ... | ... |
    | 11:00-11:30 | Pausa | | | |          sin objetivo ni método: fila de pausa
    Producto del día: ...                  o "Produit du jour" / "Product of the day", en negrita
    Otras líneas: notas del día
    ## Calendario del proyecto             segunda sección "## " opcional: tabla de
    | Mes | Fechas | Fase | Actividades | Responsables | (Objetivo) |   5 o 6 columnas

Config (JSON):
    {
      "actividad": "YEX01. Título de la actividad",
      "organizaciones": "Organización A (ES, solicitante), Organización B (FR)...",
      "duracion": "8 días de actividad (19 a 26/04/2027) y 2 de viaje...",
      "ciudad": "París. Albergue ...",
      "pais": "Francia",
      "inicio": "19/04/2027",
      "fin": "26/04/2027",
      "idioma": "es",                           es, fr o en (etiquetas "Objetivo", "Conduce")
      "titulo_calendario": "CALENDARIO DEL PROYECTO (12 MESES)"   opcional
    }

La plantilla por defecto está en proyectos/referencias/plantillas/ (fuera del
repositorio). En la hoja "Youth Exchanges" se rellenan las casillas de cabecera,
se borran los tres días de ejemplo y se escribe un bloque por día, con la franja
AM o PM y la hora en la columna A, sesión, objetivo y quién conduce en B:E y el
método en F:J. El calendario va a una hoja nueva. La hoja "Preparatory Visits"
se borra salvo con --conservar-visitas.
"""
import argparse
import copy
import json
import math
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

# Raíz del repositorio: assets -> redaccion-proyectos -> skills -> .claude -> raíz
RAIZ = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", ".."))
PLANTILLA = os.path.join(RAIZ, "proyectos", "referencias", "plantillas", "KA152_Annex_Timetable.xlsx")

T = {
    "es": dict(obj="Objetivo", cond="Conduce", cal="Calendario del proyecto", titulo_cal="CALENDARIO DEL PROYECTO"),
    "fr": dict(obj="Objectif", cond="Animation", cal="Calendrier du projet", titulo_cal="CALENDRIER DU PROJET"),
    "en": dict(obj="Aim", cond="Led by", cal="Project calendar", titulo_cal="PROJECT CALENDAR"),
}
RE_PRODUCTO = re.compile(r"^\**\s*(producto del d[ií]a|produit du jour|product of the day)", re.I)
RE_BR = re.compile(r"<br\s*/?>", re.I)
ALTO_LINEA = 13.5  # puntos por línea de texto de 10 pt


def limpio(t):
    t = RE_BR.sub("\n", t)
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", t)
    return t.strip()


def celdas(linea):
    partes = re.split(r"(?<!\\)\|", linea.strip())
    if partes and not partes[0].strip():
        partes = partes[1:]
    if partes and not partes[-1].strip():
        partes = partes[:-1]
    return [p.strip() for p in partes]


def tabla(lineas):
    """Devuelve (cabecera, filas) de las líneas de una tabla markdown."""
    filas = [celdas(l) for l in lineas if l.strip().startswith("|")]
    filas = [f for f in filas if not (f and all(re.match(r"^:?-+:?$", c) for c in f))]
    return (filas[0], filas[1:]) if filas else ([], [])


def leer_md(ruta):
    with open(ruta, encoding="utf-8-sig") as fh:
        md = fh.read().replace("\r\n", "\n")
    secciones = re.split(r"(?m)^## ", md)[1:]
    if not secciones:
        raise SystemExit("El markdown no tiene ninguna sección '## '")
    dias = []
    for bloque in re.split(r"(?m)^### ", secciones[0])[1:]:
        cab, _, resto = bloque.partition("\n")
        lineas = resto.split("\n")
        prosa = [l.strip() for l in lineas if l.strip() and not l.strip().startswith("|")
                 and not re.match(r"^-{3,}$", l.strip())]
        producto = [p for p in prosa if RE_PRODUCTO.match(p)]
        resto_prosa = [p for p in prosa if not RE_PRODUCTO.match(p)]
        objetivo = resto_prosa[0] if resto_prosa else ""
        dias.append(dict(cab=cab.strip(), objetivo=objetivo, filas=tabla(lineas)[1],
                         producto=producto, notas=resto_prosa[1:]))
    cal_cab, cal = [], []
    if len(secciones) > 1:
        cal_cab, cal = tabla(secciones[1].split("\n"))
    return dias, cal_cab, cal


def franja(hora):
    m = re.match(r"\s*(\d{1,2})\s*[:h.]\s*(\d{2})?", hora)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2) or 0)
    return "AM" if (h, mi) < (13, 30) else "PM"


def lineas_texto(texto, ancho_chars):
    return sum(max(1, math.ceil(len(p) / max(ancho_chars, 1))) for p in texto.split("\n"))


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass
    ap = argparse.ArgumentParser(description="Horario en markdown a la plantilla oficial KA152 (xlsx).")
    ap.add_argument("--md", required=True)
    ap.add_argument("--config", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--plantilla", default=PLANTILLA)
    ap.add_argument("--conservar-visitas", action="store_true", help="no borra la hoja Preparatory Visits")
    a = ap.parse_args(argv)

    if not os.path.exists(a.plantilla):
        print("AVISO: no encuentro la plantilla oficial en %s. Descárgala del portal (Annex Timetable "
              "KA152) y guárdala ahí o pásala con --plantilla." % a.plantilla, file=sys.stderr)
        return 2
    with open(a.config, encoding="utf-8-sig") as fh:
        cfg = json.load(fh)
    t = T.get(cfg.get("idioma", "es"), T["es"])
    dias, cal_cab, cal = leer_md(a.md)

    wb = openpyxl.load_workbook(a.plantilla)
    if not a.conservar_visitas:
        for nombre in list(wb.sheetnames):
            if nombre.strip().lower().startswith("preparatory"):
                del wb[nombre]
    ws = wb["Youth Exchanges"] if "Youth Exchanges" in wb.sheetnames else wb.worksheets[0]
    st = {k: copy.copy(ws[k]._style) for k in ("A1", "A9", "B9", "A10", "A11", "B11", "F11")}
    vacio = copy.copy(ws.cell(26, 11)._style)
    alto_min = ws.row_dimensions[11].height or 25

    # Anchura útil en caracteres de 10 pt (1 unidad de columna ~ 1,1 caracteres)
    ancho = {c: (ws.column_dimensions[c].width or 13) for c in "ABCDEFGHIJ"}
    ch_be = int(sum(ancho[c] for c in "BCDE") * 1.1) - 2
    ch_fj = int(sum(ancho[c] for c in "FGHIJ") * 1.1) - 2
    ch_aj = int(sum(ancho.values()) * 1.1) - 2

    izq = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["B2"] = cfg.get("actividad", "")
    ws["B3"] = cfg.get("organizaciones", "")
    ws["B4"] = cfg.get("duracion", "")
    ws["A7"] = cfg.get("ciudad", "")
    ws["D7"] = cfg.get("pais", "")
    ws["G7"] = cfg.get("inicio", "")
    ws["I7"] = cfg.get("fin", "")
    for c in ("B2", "B3", "B4"):
        ws[c].alignment = izq
    ws["A7"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila, celda, chars in ((3, "B3", ch_aj - 12), (4, "B4", ch_aj - 12)):
        ws.row_dimensions[fila].height = max(ws.row_dimensions[fila].height or 30,
                                             ALTO_LINEA * lineas_texto(str(ws[celda].value), chars) + 8)
    ws.row_dimensions[7].height = max(28, ALTO_LINEA * lineas_texto(str(ws["A7"].value), 40) + 8)

    # Borrar los días de ejemplo (filas 10 en adelante)
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 10:
            ws.unmerge_cells(str(rng))
    for r in range(10, ws.max_row + 1):
        ws.row_dimensions[r].height = None
        for c in range(1, 11):
            ws.cell(r, c).value = None
            ws.cell(r, c)._style = copy.copy(vacio)

    def fila_ancha(r, texto, estilo, fuente):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        ws.cell(r, 1).value = texto
        ws.cell(r, 1)._style = copy.copy(estilo)
        for c in range(2, 11):
            ws.cell(r, c)._style = copy.copy(estilo)
        ws.cell(r, 1).alignment = izq
        ws.cell(r, 1).font = fuente
        ws.row_dimensions[r].height = max(18, ALTO_LINEA * lineas_texto(texto, ch_aj) + 6)

    r = 10
    n_filas = 0
    for d in dias:
        fila_ancha(r, d["cab"].upper(), st["A10"], Font(bold=True, size=11))
        ws.cell(r, 1).alignment = Alignment(horizontal="left", vertical="center")
        r += 1
        if d["objetivo"]:
            fila_ancha(r, limpio(d["objetivo"]), st["B11"], Font(italic=True, size=10))
            r += 1
        for f in d["filas"]:
            hora, ses, obj, met, cond = (limpio(x) for x in (f + [""] * 5)[:5])
            pausa = not obj and not met
            fr = franja(hora)
            ws.cell(r, 1).value = "%s\n%s" % (fr, hora) if fr else hora
            ws.cell(r, 1)._style = copy.copy(st["A11"])
            ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=10)
            for c in range(2, 11):
                ws.cell(r, c)._style = copy.copy(st["B11"] if c < 6 else st["F11"])
            texto_b = ses if pausa else "%s\n%s: %s%s" % (ses, t["obj"], obj,
                                                          "\n%s: %s" % (t["cond"], cond) if cond else "")
            ws.cell(r, 2).value = texto_b
            ws.cell(r, 6).value = met or None
            ws.cell(r, 2).alignment = izq
            ws.cell(r, 6).alignment = izq
            ws.cell(r, 2).font = Font(size=10, bold=pausa, italic=pausa)
            ws.cell(r, 6).font = Font(size=10)
            lineas = max(lineas_texto(texto_b, ch_be), lineas_texto(met, ch_fj), lineas_texto(hora, 12) + 1)
            ws.row_dimensions[r].height = alto_min if pausa else max(alto_min, ALTO_LINEA * lineas + 6)
            r += 1
            n_filas += 1
        for extra in d["producto"] + d["notas"]:
            fila_ancha(r, limpio(extra), st["B11"], Font(bold=extra in d["producto"], size=10))
            r += 1

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "9:9"

    if cal_cab:
        wc = wb.create_sheet(cfg.get("hoja_calendario", t["cal"])[:31])
        n = len(cal_cab)
        anchos = {5: [10, 24, 18, 70, 46], 6: [10, 16, 18, 60, 36, 36]}.get(n)
        if not anchos:
            anchos = [10, 16, 18] + [max(24, int(130 / max(1, n - 3)))] * (n - 3) if n > 3 else [30] * n
        for i, w in enumerate(anchos):
            wc.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        wc.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        wc.cell(1, 1).value = cfg.get("titulo_calendario", t["titulo_cal"])
        wc.cell(1, 1)._style = copy.copy(st["A1"])
        wc.row_dimensions[1].height = 24
        for j, h in enumerate(cal_cab, start=1):
            c = wc.cell(3, j)
            c.value = limpio(h)
            c._style = copy.copy(st["B9"])
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wc.row_dimensions[3].height = 32
        for i, fila in enumerate(cal, start=4):
            fila = (fila + [""] * n)[:n]
            for j, v in enumerate(fila, start=1):
                c = wc.cell(i, j)
                c.value = limpio(v)
                c._style = copy.copy(st["B11"])
                c.alignment = Alignment(horizontal="left" if j > 3 else "center", vertical="center", wrap_text=True)
                c.font = Font(size=10, bold=j <= 3)
            lineas = max(lineas_texto(limpio(v), int(anchos[j] * 1.1) - 2) for j, v in enumerate(fila))
            wc.row_dimensions[i].height = max(30, ALTO_LINEA * lineas + 8)
        wc.page_setup.orientation = "landscape"
        wc.page_setup.fitToWidth = 1
        wc.page_setup.fitToHeight = 0
        wc.sheet_properties.pageSetUpPr.fitToPage = True
        wc.print_title_rows = "3:3"

    wb.active = 0
    wb.save(a.out)
    print("OK %s: %d días, %d filas de sesiones, calendario de %d filas y %d columnas"
          % (a.out, len(dias), n_filas, len(cal), len(cal_cab)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
